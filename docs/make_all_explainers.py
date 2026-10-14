# ruff: noqa: E501
"""Builds one plain-language Word guide (with an architecture section) per project, plus an overview.

Run: python docs/make_all_explainers.py     Output: C:/PROJECTS/project_explainers/
"""

import sys
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt

OUT = Path(sys.argv[1] if len(sys.argv) > 1 else r"C:\PROJECTS\project_explainers")

PROJECTS = [
    dict(
        file="EvalKit", name="EvalKit", tagline="A test machine that tells you, with numbers, whether your AI app is actually good",
        path="C:\\PROJECTS\\evalkit",
        simple="You build an AI app and it gives answers. How do you know the answers are good? EvalKit asks your app a list of questions with known right answers, grades every reply, and shows a score, the failures, and whether your latest change helped or hurt. Think of school exams for your AI.",
        why=["Changes to prompts or models often break things silently; EvalKit catches that before users do.",
             "Compares two versions with a statistical test, so luck is not mistaken for improvement.",
             "Plugs into any app in minutes, with no code written per app.",
             "Shows employers you know how to measure AI quality, the hard part of AI engineering."],
        features=["Connect any app: paste a cURL command, pick a service (OpenAI, Claude, Gemini, Ollama, Groq), or give a web address",
                  "Handles streaming replies, file uploads, websites with a login, and Python functions",
                  "Scorers: exact, contains, pattern, similarity, faithfulness, abstain, latency, cost, AI judge (1 to 5)",
                  "Search-quality scores: hit rate, MRR, recall", "A/B comparison, consistency checks, red-team attacks, test generation",
                  "CI quality gate, web dashboard with live progress, safe database copies (snapshot)"],
        stack="Python, FastAPI, SQLite, Typer CLI, React + TypeScript + Tailwind, Docker, Gemini API (optional)",
        arch="""Your app (any API / website / Python function)
        ^   question            |  reply
        |                       v
  +-----------+   +----------------------------+
  |  Targets  |-->|  Runner (parallel, repeats) |
  | http/web/ |   +-------------+--------------+
  |  python   |                 v
  +-----------+        Scorers (plug-ins) + LLM judge
                                |
                                v
   SQLite storage --> Summary / A-B stats / Quality gate
                                |
                  +-------------+-------------+
                  v                           v
          CLI (evalkit ...)        FastAPI + React dashboard""",
        flow=["Connect layer reads your app's reply and finds the answer, passages and token counts on its own.",
              "Runner asks every question in parallel and keeps each reply.", "Each scorer plug-in grades a reply; results are saved in SQLite.",
              "Summary adds pass rate, confidence range, speed and cost; compare tests runs A vs B; the gate returns pass or fail.",
              "The dashboard and CLI both read the same database."],
        run=["docker compose up -d --build, then open http://localhost:8800", "Optional Gemini key goes in .env (never in code)"],
        limits=["Built for question-and-answer and document-search apps, not image or video apps.",
                "The AI judge needs an LLM key; the free Gemini tier allows about 20 requests a day per model.",
                "Website-login support was tested on a simulated site, not a real one yet."],
        pitch="I built a toolkit that plugs into any AI app, runs a question set, scores the answers and tells you statistically whether a change made things better or worse. It can block a bad release in CI, and I used it on my own apps to find real bugs."),
    dict(
        file="Tapedeck", name="Tapedeck", tagline="Record an API's traffic once, then mock it, replay it and load-test it",
        path="C:\\PROJECTS\\tapedeck",
        simple="Like a voice recorder for the conversation between an app and its server. Once you have a recording (a tape) you can play the server's answers back as a fake server, replay the requests against a new server version to see what changed, or send them as a crowd to see when the server struggles.",
        why=["Frontend work does not stop when the backend is down or unfinished (mock).",
             "Shows exactly which fields changed between two server versions (replay and diff).",
             "Shows how many users a server can take before slowing down (load test).",
             "Secrets such as tokens and passwords are hidden before a tape is saved."],
        features=["Import a HAR file or record live through a proxy", "Smart request matcher and noise detection", "Mock server from a tape",
                  "Replay with side-by-side response diff", "Load test with live charts over WebSocket", "Web UI, CLI, Docker, benchmarks"],
        stack="Python, FastAPI, SQLite (one file per tape), React, WebSocket, Docker",
        arch="""  App / Browser                       Real server
       |                                   ^
       v          record                   |
  +---------+  (secrets scrubbed)   +------+-----+
  |  Proxy  |---------------------->|   Tape     |
  +---------+                       |  (SQLite)  |
                                    +--+--+--+---+
                          +------------+  |  +-------------+
                          v               v                v
                    Mock server       Replay + diff     Load test
                  (fake backend)    (new vs old server)  (live chart)
                          \\______________|________________/
                                         v
                              Web UI (:8700) and CLI""",
        flow=["Capture traffic from a HAR file or through the recording proxy.", "Scrub tokens, cookies and passwords, then save as a tape.",
              "Pick a mode: mock, replay or load test.", "Results show in the web page, with live charts for load tests."],
        run=["tapedeck ui opens http://localhost:8700", "Or: docker compose up -d --build"],
        limits=["Works with HTTP APIs; it is not a general network tool.", "Not yet in a git repository or published."],
        pitch="Record an API once, then use that recording as a fake server, a regression checker and a load tester."),
    dict(
        file="Apexcut", name="Apexcut", tagline="Turns a long racing video into highlight clips and shows how it decided",
        path="C:\\PROJECTS\\Apexcut", repo="https://github.com/hridasha/Apexcut",
        simple="Give it a YouTube link or video file. Four AI models read every frame, Whisper and Gemini listen to the commentary, and the app cuts the best moments into clips you can trim, caption, add music to, and export as one video. You can watch the footage with every model's output drawn on it.",
        why=["Saves hours of manual video editing for racing content.", "Not a black box: you can see what each model detected and why a moment was picked.",
             "Runs fully on your own machine with Docker; no cloud storage needed."],
        features=["Event detection (YOLO), segmentation of wall, curb, track and car", "Camera-view classifier", "Speech-to-text with Whisper plus Gemini",
                  "Clip editor: trim, captions, music, stickers, export", "Live annotated view streamed as MJPEG", "Accounts, plans, and push notifications when processing is done"],
        stack="React 19 + TypeScript + Tailwind + Vite, FastAPI, Tortoise ORM, Postgres, PyTorch, TensorFlow, Whisper, MoviePy, yt-dlp, Docker",
        arch="""  Browser (React, nginx :5173)
     |  REST + JWT                      ^ "processing done"
     v                                  |
  API :5001 (FastAPI, Tortoise) --> Socket relay :8005 (WebSocket)
     |    \\                                 ^
     |     +--> Postgres :5433              | notifications
     v
  Worker :5002 (FastAPI + PyTorch, TensorFlow,
                Whisper, MoviePy, yt-dlp)
     +--> shared ./storage  <-- served by API at /media""",
        flow=["You submit a link or file; the API stores it and hands the job to the worker.", "The worker runs the models on every frame and transcribes the audio.",
              "Exciting moments become clips; the live annotated view streams while it works.", "The socket relay tells your browser when it is done; you edit and export."],
        run=["docker compose up (see the README in the project)", "Optional API keys for Gemini"],
        limits=["API and worker share folders, so they must run on the same machine.", "Heavy models need a decent computer."],
        pitch="An app that watches a whole race and cuts the highlights for you, while showing exactly how each AI model made its decision."),
    dict(
        file="EchoMind", name="EchoMind", tagline="Chat with an AI version of a person built from their memories",
        path="C:\\PROJECTS\\Echomind",
        simple="You give it a person's details, a photo and anything that captures them (letters, documents, voice notes, videos). You can then chat with an AI version that answers from those memories, shows an expression matching each reply, and can speak, optionally as a lip-synced video.",
        why=["Keeps stories and voices of loved ones in a form you can talk to.", "Answers come from the supplied memories, not made up from nothing.",
             "Only one API key is required; every other service is an optional switch."],
        features=["Upload photos, letters, documents, voice notes, videos", "Memory search with vectors (pgvector)", "Emotion shown with each reply", "Voice with ElevenLabs or edge-tts",
                  "Optional talking-face video (D-ID or MuseTalk)", "Web search tool per chat (MCP + DuckDuckGo)"],
        stack="FastAPI, Postgres + pgvector, RabbitMQ, worker, nginx, Gemini API, Docker",
        arch="""Browser --:8080 nginx (page + /files)--/backend--> FastAPI :8001
                                                       |
     +-----------+-----------+------------+------------+---------+
     v           v           v            v            v         v
 Postgres    RabbitMQ    shared       Gemini API   MCP web   ElevenLabs /
 + pgvector     |        volume       (chat,       search    edge-tts
 (users,        v        (photos,     embeddings,
 chats,      worker      clips,       speech,
 memories)  (read files, audio,       captions)
            save memory, video)       D-ID or MuseTalk :8002
            clips, voice)             (optional talking face)""",
        flow=["Files are uploaded; the API queues the work in RabbitMQ.", "The worker reads files, saves memories as vectors, and makes emotion clips.",
              "When you chat, the API finds the closest memories and asks Gemini to answer from them.", "The reply gets an expression and optional voice or video."],
        run=["Needs Docker Desktop and one Gemini key in .env", "See the README quick start"],
        limits=["Many services must run together.", "Quality depends on how much material you provide."],
        pitch="A chat with an AI version of someone, answering only from their real memories."),
    dict(
        file="PaperGraph", name="PaperGraph", tagline="Turn research papers into a knowledge graph you can chat with",
        path="C:\\PROJECTS\\PaperGraph",
        simple="Upload PDFs or fetch from arXiv. An LLM pulls out facts like (subject, relation, object), merges them into one graph of concepts, and you ask questions that are answered by walking through that graph.",
        why=["See how ideas across several papers connect.", "Answers show the path they came from, so you can check them.", "Finds clusters of related concepts and surprising gaps."],
        features=["Ingest PDFs or arXiv links with live progress", "Interactive force graph (D3)", "Chat with history and path highlighting", "Per-paper counts and removal",
                  "Search, export JSON or CSV", "Community detection (Louvain) and surprise detection"],
        stack="FastAPI, NetworkX, D3.js, Gemini, WebSocket, uv",
        arch="""  PDF / arXiv URL
        |
        v
  Ingestion --> LLM triple extraction --> Graph (NetworkX)
                                              |
        +-------------------------------------+-----------+
        v                      v                            v
   Node embeddings      Insights (clusters,        Export (JSON/CSV)
        |               surprises)
        v
  Question --> find related nodes --> traverse subgraph --> LLM answer
                                           |
                                    highlighted path in D3 UI""",
        flow=["Papers are parsed and facts extracted.", "Facts merge into one graph, stored on disk.", "A question is matched to graph nodes; the nearby subgraph is read.", "The LLM writes the answer and the UI highlights the path."],
        run=["uv sync, then uv run uvicorn backend.main:app --reload"],
        limits=["Answers fall back to graph-only text when Gemini is unavailable.", "In our EvalKit test it answered an out-of-scope question with unrelated graph text."],
        pitch="Ask questions across many papers and see exactly which path through the knowledge graph produced the answer."),
    dict(
        file="Incident_Copilot", name="Incident Copilot", tagline="Chat with analyzed system logs; a human approves anything flagged as an incident",
        path="C:\\PROJECTS\\incident-copilot",
        simple="An on-call engineer asks questions about a system-log dataset. The agent retrieves logs, judges how serious they look, and if serious, pauses and asks a human to approve before recording an incident.",
        why=["AI should not decide alone that something is a real incident.", "The paused state survives a server restart.", "Gives a fast way to question thousands of log lines."],
        features=["LangGraph agent: retrieve, triage, route", "Human-in-the-loop approval via interrupt()", "Durable checkpointer (SQLite)", "Optional API-key header", "Chat web page"],
        stack="LangGraph, FastAPI, Gemini embeddings and chat, SQLite checkpointer",
        arch="""question --> retrieve --> triage --> route
                 |           |          |
          log index     LLM judges   serious? --yes--> PAUSE (interrupt)
        (embeddings)    severity            |              |
                                            no       human approves/dismisses
                                            |              |
                                            v              v
                                         answer      record incident
        State saved in a checkpointer (SQLite) per thread_id""",
        flow=["Question arrives with a thread id.", "Relevant log lines are retrieved by meaning.", "The agent decides whether it looks serious.", "If serious it pauses; /api/confirm resumes with approve or dismiss."],
        run=["python -m uvicorn backend.main:app --reload (inside its venv)"],
        limits=["Depends on the analyzed log file and a Gemini key."],
        pitch="An incident chatbot that asks a human before it declares anything a real incident."),
    dict(
        file="FinQA_Agent", name="FinQA Agent", tagline="Finds the right financial-filing page for a question",
        path="C:\\PROJECTS\\finqa-agent",
        simple="It holds 299 pages from financial filings. Given a question it finds the best matching pages by meaning, so a later step can calculate the answer.",
        why=["Shows retrieval working on real financial data.", "Has its own retrieval self-check endpoint.", "Embeddings are cached, so asking is quick."],
        features=["Build a vector index of filing excerpts", "Ask: top-k best excerpts with scores", "Retrieval evaluation endpoint", "Session history reset", "Small web page"],
        stack="FastAPI, Gemini embeddings, NumPy, FinQA dataset",
        arch="""dev.json (FinQA) --> loader --> filing excerpts
                                      |
                                      v
                             embeddings (cached)
                                      |
 question --> embed --> similarity search (top-k) --> excerpts + scores
                                      |
                              /api/ask  /api/eval_retrieval""",
        flow=["Excerpts are embedded once and cached.", "A question is embedded.", "The closest excerpts are returned with scores."],
        run=["python -m uvicorn backend.main:app (inside its venv)"],
        limits=["The ask endpoint only retrieves; it does not write a final answer.", "A new question needs one embedding call."],
        pitch="Retrieval over real financial filings; EvalKit measured hit@5 0.92 and MRR 0.86 on it."),
    dict(
        file="RAG_Chatbot", name="RAG Chatbot", tagline="Chat with your own PDFs and CSV files",
        path="C:\\PROJECTS\\rag-chatbot",
        simple="Upload a document or a spreadsheet and ask questions. For documents it finds the relevant pieces and answers from them; for CSVs it can build a database query; otherwise it falls back to general knowledge. Replies stream in live.",
        why=["Answers come from your own files.", "Streaming replies feel fast.", "Rotates between several API keys when one is rate-limited."],
        features=["PDF and CSV upload", "Chunks and vector store (Chroma)", "CSV to SQL answers", "Streaming answers", "Usage counter per key", "Plain HTML frontend"],
        stack="FastAPI, Gemini, ChromaDB, SQLite, HTML/JS",
        arch="""Browser --message--> /chat (FastAPI, streams SSE)
                                  |
          +-----------------------+------------------+
          v                       v                  v
   Document question       CSV question        General question
   Chroma search           LLM writes SQL      straight to Gemini
          |                   |    |                  |
          v                   v    v                  v
   chunks + prompt        SQLite result       streamed answer
          \\__________________|__________________/
                             v
                 Gemini (key rotation) --> streamed deltas""",
        flow=["You upload a file; it is chunked and stored.", "A message decides the route: document, CSV or general.", "Gemini writes the answer, streamed word by word."],
        run=["Create venv, install requirements, add Gemini key in .env, python -m backend.main"],
        limits=["When all Gemini keys hit their limit, chat returns an error (seen during our EvalKit test)."],
        pitch="A chatbot grounded in your own documents and spreadsheets, with streaming answers."),
    dict(
        file="Data_Insights_Studio", name="Data Insights Studio", tagline="Upload a dataset, get profiling, charts, statistics, AutoML and AI insights",
        path="C:\\PROJECTS\\data-insights-studio",
        simple="Upload a dataset (or use the sample). It profiles it automatically, lets you build charts, runs statistical tests, trains simple ML models, and lets you ask questions about the data in plain English.",
        why=["One place for the first look at any dataset.", "A client intake form recommends a right-sized solution and avoids over-engineering.", "Questions in plain English instead of code."],
        features=["Overview: schema, missing values, duplicates, outliers", "Visual explorer with many chart types", "Statistics: correlation, group comparison, chi-square", "AutoML: regression, classification, clustering",
                  "AI insights and Q&A (Gemini)", "Solution Builder intake wizard"],
        stack="FastAPI, pandas, scikit-learn, scipy, React + TypeScript + Tailwind + Recharts, Gemini",
        arch="""Browser (React + Recharts)
      |  REST
      v
  FastAPI backend
      +--> dataset state (pandas) --> profile / outliers
      +--> charts, stats (scipy), AutoML (scikit-learn)
      +--> Gemini: insights, Q&A, Solution Builder
      +--> optional Colab worker for training (/train)""",
        flow=["Upload or load the sample.", "Profile and outlier scan run.", "Explore with charts and tests, train models.", "Ask questions or use the Solution Builder."],
        run=["pip install -r backend/requirements.txt, uvicorn backend.main:app, npm run dev in frontend"],
        limits=["No virtual environment was set up here, so it was not run in the EvalKit test."],
        pitch="Drop in data and get a full analysis workspace with AI explanations."),
    dict(
        file="Browser_Automation_Agent", name="Browser Automation Agent", tagline="Give it a goal in English; it searches the web and returns structured data",
        path="C:\\PROJECTS\\browser_automation_project",
        simple="You say, for example, find remote AI internships. A chain of AI steps plans a search, finds pages, visits them with a real headless browser, pulls out structured fields with an LLM, and checks the result quality, retrying if needed.",
        why=["Automates repetitive web research.", "Returns tidy structured data, not a pile of pages.", "Available as CLI, REST API with live streaming, and an MCP server."],
        features=["Planner, search, browser, extraction, validator steps", "Parallel page visits with Playwright", "Retries up to 3 times with a reformulated query", "CLI, REST + SSE dashboard, MCP server"],
        stack="LangGraph, Tavily, Playwright, Groq (Llama 3.3), FastAPI, uv",
        arch="""planner --> search --> browser --> extraction --> validator --done--> result
                ^                                          |
                +------------- retry (up to 3x) -----------+
 planner: goal -> query + fields     search: Tavily URLs
 browser: Playwright, one tab per URL   extraction: Groq / Llama 3.3
 Interfaces: CLI (main.py) | REST + SSE (api/server.py) | MCP (stdio)""",
        flow=["Planner turns your goal into a search query and field list.", "Search returns candidate URLs.", "Browser visits all pages in parallel.", "LLM extracts fields; validator drops weak results or loops back."],
        run=["uv sync, uv run playwright install chromium, copy .env.example to .env, run main.py"],
        limits=["Needs Tavily and Groq keys.", "Result quality depends on the pages found."],
        pitch="An agent that browses the web for you and hands back clean, structured data."),
    dict(
        file="Log_Pattern_Agent", name="Log Pattern Agent", tagline="Reads system logs, finds anomalies and groups them into patterns",
        path="C:\\PROJECTS\\log-pattern-agent",
        simple="It reads log lines (from the BGL supercomputer log dataset), asks an LLM to analyze each line, stores the findings, then clusters similar findings into patterns and prints a report. It remembers what it already analyzed, so re-runs only handle new lines.",
        why=["Turns thousands of raw log lines into a few readable patterns.", "Saves effort by never re-analyzing the same line.", "Its analyzed output feeds Incident Copilot."],
        features=["BGL log loader", "LLM analysis per line (severity, category, anomaly flag, reasoning)", "Embeddings and clustering into patterns", "Persistent memory file", "Evaluation script"],
        stack="Python, Gemini (analysis + embeddings), JSON store",
        arch="""BGL log --> loader --> new lines only
                               |
                               v
                  analyzer (LLM) --> findings
                               |
                    store (analyzed_log.json, memory)
                               |
                  embeddings --> pattern matcher (clusters)
                               |
                               v
                            report""",
        flow=["Load log lines.", "Skip lines already analyzed.", "Analyze new ones and store the results.", "Cluster all findings into patterns and report."],
        run=["python run.py (see arguments in run.py)"],
        limits=["Command-line tool without a web page.", "Needs a Gemini key."],
        pitch="Summarizes huge logs into a short list of recurring problem patterns."),
    dict(
        file="Screenshot_to_Draft", name="Screenshot to Gmail Draft", tagline="Upload a LinkedIn or job screenshot; get a follow-up email saved as a Gmail draft",
        path="C:\\PROJECTS\\screenshot-to-draft",
        simple="You upload a screenshot of a profile or job application. Gemini reads the contact details and writes a follow-up email, and the app saves it as a Gmail draft. It never sends the email; you review and send it yourself.",
        why=["Saves time on networking and job follow-ups.", "Safe by design: draft only, never sends.", "Free to run with a free Gemini key."],
        features=["Screenshot reading with Gemini vision", "Drafts a follow-up email", "Saves to Gmail through OAuth", "Simple web page"],
        stack="Python, Flask, Gemini, Gmail API",
        arch="""Browser --upload screenshot--> app.py
                                 |
                                 v
                      Gemini (vision): extract name, contact, role
                                 |
                      graph.py: write follow-up email
                                 |
                      gmail_service: create DRAFT (never send)
                                 |
                      You review and send in Gmail""",
        flow=["Upload the screenshot.", "Gemini extracts details.", "An email is drafted.", "It is saved to Gmail drafts for you."],
        run=["pip install -r requirements.txt, add Gemini key, set up Gmail OAuth once, run app.py"],
        limits=["Gmail OAuth setup is a one-time manual step."],
        pitch="Turn a screenshot into a ready-to-review follow-up email, without ever auto-sending."),
    dict(
        file="Tender_Watcher", name="Tender Watcher", tagline="Watches tender portals daily and alerts on new tenders",
        path="C:\\PROJECTS\\tender-watcher",
        simple="Users add tender portals to watch. A daily check scrapes them, spots items it has not seen before, and sends alerts through a Telegram bot. Results can be exported to Excel. It is designed to run on a phone-friendly website.",
        why=["Nobody has to check many portals by hand.", "Alerts arrive where people already are (Telegram).", "Deployed free-tier friendly (Render, Neon, GitHub Actions cron)."],
        features=["Portal list with a large starter CSV", "Scraper plus a JavaScript-render fallback", "Seen-item tracking", "Telegram bot alerts", "Excel export", "Users and security", "Manual Check now button"],
        stack="FastAPI, SQLAlchemy, Jinja2 templates, Postgres (Neon), Telegram bot, Render, GitHub Actions",
        arch="""GitHub Actions cron (daily) --> /api/run-check
Manual "Check now" button ----------^
                |
                v
        checker.py --> scraper (static or JS-rendered pages)
                |
                v
        compare with SeenItem table (Postgres/SQLite)
                |
                +--> new items --> Telegram bot --> user's phone
                +--> web dashboard / Excel export (FastAPI + Jinja2)""",
        flow=["The daily trigger starts the check.", "Each portal is scraped.", "New items are compared with ones already seen.", "New ones are sent through Telegram and shown on the site."],
        run=["See DEPLOY.md for Neon, Telegram bot and Render steps"],
        limits=["Scrapers can break when a portal changes its layout."],
        pitch="A daily watcher that tells you about new tenders so you never miss one."),
    dict(
        file="Prescription_Reader", name="Prescription Reader", tagline="Reads a doctor's handwritten medicine name from an image",
        path="C:\\PROJECTS\\Doctors-handwritten-prescription-recognition",
        simple="Upload a cropped image of a handwritten medicine name. A MobileNetV2 classifier trained on 78 medicine names tells you what it says, how confident it is, and the generic name.",
        why=["Handwriting is hard to read; this gives a ranked guess with confidence.", "Shows several models side by side with measured accuracy.", "A notebook project turned into a usable web app."],
        features=["Single and batch reading", "Top guess plus next closest guesses", "Compare 9 trained models", "Prediction history with export", "Searchable medicine list"],
        stack="TensorFlow 2.15 (MobileNetV2), FastAPI, Python 3.11, uv, browser UI",
        arch="""Browser (upload word image)
      |
      v
  FastAPI backend
      +--> preprocess image
      +--> MobileNetV2 (one of 9 trained models)
      +--> top-k classes + confidence + generic name
      +--> history log (thumbnails, JSON export)""",
        flow=["Upload an image.", "It is resized and normalized.", "The model predicts one of 78 names.", "The UI shows the guess, alternatives and generic name."],
        run=["uv sync, uv run uvicorn backend.main:app --reload"],
        limits=["Reads single cropped words only, and only the 78 known medicines.", "Not for medical use; a pharmacist must confirm."],
        pitch="Computer vision that reads doctor handwriting with a confidence score."),
    dict(
        file="Drowsiness_Detection", name="Drowsiness Detection", tagline="Warns you in real time when your eyes close too long or you yawn",
        path="C:\\PROJECTS\\Drowsiness-detection",
        simple="It watches your webcam, tracks points on your eyes and mouth, and raises a warning and then an alarm if your eyes stay closed or you yawn repeatedly.",
        why=["Could help prevent accidents from tiredness.", "Needs no training data; works out of the box.", "Thresholds are adjustable live."],
        features=["Live annotated video", "Two alert levels with sound", "Separate yawn detection", "Live sensitivity settings saved to disk", "Session history with durations and severity"],
        stack="MediaPipe FaceMesh, FastAPI, browser dashboard, Python 3.11",
        arch="""Webcam --> browser dashboard (Start Monitoring)
                     |
                     v
              FastAPI backend
                     +--> MediaPipe FaceMesh landmarks
                     +--> EAR (eye) / MAR (mouth) math
                     +--> frame counters --> warning / alarm tiers
                     +--> episodes log + saved settings
                     v
        annotated video stream + alerts to the browser""",
        flow=["Frames come from the camera.", "Landmarks give eye and mouth ratios.", "Several closed-eye frames in a row raise warning, then alarm.", "Episodes are logged."],
        run=["uv sync, uv run uvicorn backend.main:app --reload, open http://127.0.0.1:8000"],
        limits=["Lighting, glasses and camera angle affect accuracy.", "A helper tool, not a certified safety device."],
        pitch="A real-time tiredness alarm built from face landmarks and simple math."),
    dict(
        file="MoodyTunes", name="MoodyTunes", tagline="A music player that picks songs for your mood",
        path="C:\\PROJECTS\\MoodyTunes",
        simple="A website where you pick, type or show your mood (through the webcam) and it plays matching songs. It also has playlists you can share, favourites, ratings, and Spotify search.",
        why=["Music matched to how you feel.", "Mood history shows patterns over time.", "Collaborative playlists and an admin review queue for uploads."],
        features=["Mood by picking, webcam (DeepFace) or typed text (HuggingFace model)", "Songs by language, mood, artist", "Collaborative playlists", "Favourites, ratings, popular page, recently played",
                  "Mood history heatmap", "User uploads with admin approval", "Spotify connection"],
        stack="Django, DeepFace, HuggingFace emotion model, Spotify OAuth, SQLite",
        arch="""Browser
   |  pages + audio player
   v
 Django (auth, songs, playlists, ratings)
   +--> /cam/        one webcam frame --> DeepFace emotion
   +--> /mood-text/  text --> HuggingFace emotion model
   +--> mood history (per user) --> charts / heatmap
   +--> Spotify OAuth --> search
   +--> Django admin: approve uploaded songs
   v
 SQLite database + media files""",
        flow=["You choose or detect a mood.", "The mood is logged.", "Songs matching the mood are listed.", "The shared player plays them and changes colour with the mood."],
        run=["pip install -r requirements.txt, python manage.py migrate, python manage.py runserver"],
        limits=["Emotion models are guesses and can be wrong.", "Spotify features need your own credentials."],
        pitch="A mood-aware music player that detects emotion from your camera or your words."),
    dict(
        file="Yoga_Pose_Correction", name="Yoga Pose Correction System", tagline="Upload a photo or use the camera to get feedback on your yoga pose",
        path="C:\\PROJECTS\\Yoga-pose-correction-system",
        simple="It identifies your yoga pose with a trained model, measures joint angles with pose landmarks, and tells you where your angles differ from the ideal pose.",
        why=["Feedback when no teacher is around.", "Angle-based corrections are concrete and easy to follow.", "Works from a photo or live camera."],
        features=["Accounts and profiles", "Pose library with benefits, difficulty, hold time", "Photo upload with correction feedback", "Live webcam correction", "Keras pose classifier with MediaPipe landmarks"],
        stack="Django, FastAPI, TensorFlow/Keras, MediaPipe, Postgres (Supabase)",
        arch="""Browser --> Django (ypc): UI, auth, poses, data --> Postgres (Supabase)
                 |
                 | launches and manages
                 v
        FastAPI server (fastapi_server/): real-time webcam correction
                 +--> MediaPipe landmarks --> joint angles
                 +--> Keras classifier --> pose name
                 +--> compare with ideal angles --> feedback""",
        flow=["You choose or upload a pose.", "Landmarks and angles are measured.", "The classifier names the pose.", "Differences from ideal angles become advice."],
        run=["Create venv, pip install -r requirements.txt, copy .env.example to .env, run Django and the FastAPI service"],
        limits=["Not a replacement for a trained teacher; camera angle affects accuracy."],
        pitch="A yoga coach in your browser that points out which joint angle to fix."),
    dict(
        file="Resume_Screener", name="Resume Screener", tagline="Ranks resumes against a job and lets you ask questions about candidates",
        path="C:\\SOLGURUZ\\resume_screener_django",
        simple="A Django website that scores how well resumes match a job description, ranks candidates, and has a chatbot that answers questions about resumes with cited sources.",
        why=["Speeds up the first screening of many resumes.", "Answers come with sources so a recruiter can check them.", "Gemini extracts structured details from resumes."],
        features=["Resume matching and ranking by chunk similarity", "Chatbot with answer text and sources", "Gemini-based extraction pipeline", "Docker setup"],
        stack="Django, Gemini, embeddings, Docker (from the files inspected)",
        arch="""Recruiter --> Django website
                 +--> upload resumes --> Gemini extraction --> structured data
                 +--> matching.py: rank_by_chunk_similarity (resume vs job)
                 +--> chatbot.py: answer_question --> answer_text + sources
                 v
          Database (resumes, chunks, scores)""",
        flow=["Resumes are uploaded and parsed.", "Chunks are compared with the job description.", "Candidates are ranked.", "Questions are answered from the matching chunks."],
        run=["See the README in the project folder"],
        limits=["This guide is based on a quick look at the code, not a full run. Evaluating it with EvalKit is still to do."],
        pitch="Faster, explainable first-pass resume screening."),
    dict(
        file="Portfolio", name="Portfolio Website", tagline="Hridasha Joshi's personal portfolio site",
        path="C:\\PROJECTS\\portfolio",
        simple="A plain HTML, CSS and JavaScript website that presents projects, skills and a downloadable resume. There is no build step, so it can be hosted anywhere that serves static files.",
        why=["Gives recruiters one link to see your work.", "Free to host (GitHub Pages, Vercel or Netlify).", "Fast and simple, nothing to maintain."],
        features=["Static pages (index and resume)", "Downloadable resume PDF", "Deploy guides for GitHub Pages, Vercel, Netlify"],
        stack="HTML, CSS, JavaScript",
        arch="""Visitor's browser
     |
     v
  Static host (GitHub Pages / Vercel / Netlify)
     +--> index.html + style.css + script.js
     +--> assets/ (images, resume PDF)""",
        flow=["A visitor opens the link.", "The static files load.", "They read about projects or download the resume."],
        run=["Open index.html, or deploy the folder to any static host"],
        limits=["Update the GitHub and LinkedIn links before publishing."],
        pitch="One clean link that shows who you are and what you have built."),
]


def build(p: dict) -> Path:
    d = Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(11)
    for s in d.sections:
        s.left_margin = s.right_margin = Inches(1)

    def h(t, lvl=1):
        d.add_heading(t, lvl)

    def para(t, italic=False, bold=False):
        r = d.add_paragraph().add_run(t)
        r.italic, r.bold = italic, bold

    def bullets(items, style="List Bullet"):
        for i in items:
            d.add_paragraph(i, style=style)

    d.add_heading(p["name"], 0)
    para(p["tagline"], italic=True)
    para("Folder: " + p["path"] + (f"    Repo: {p['repo']}" if p.get("repo") else ""))
    h("1. What it is, in plain words")
    para(p["simple"])
    h("2. Why it is useful")
    bullets(p["why"])
    h("3. What it can do")
    bullets(p["features"])
    h("4. Architecture")
    para("The big picture (boxes are parts of the system, arrows show how information moves):")
    for line in p["arch"].split("\n"):
        r = d.add_paragraph().add_run(line.replace("\\\\", "\\"))
        r.font.name = "Consolas"
        r.font.size = Pt(8.5)
        d.paragraphs[-1].paragraph_format.space_after = Pt(0)
    d.add_paragraph()
    para("Technology used: " + p["stack"], bold=True)
    h("5. How it works, step by step")
    bullets(p["flow"], "List Number")
    h("6. How to run it")
    bullets(p["run"])
    h("7. Honest limits")
    bullets(p["limits"])
    h("8. How to explain it in 30 seconds")
    para('"' + p["pitch"] + '"', italic=True)
    OUT.mkdir(parents=True, exist_ok=True)
    out = OUT / f"{p['file']}_Explained.docx"
    d.save(out)
    return out


def overview() -> Path:
    d = Document()
    d.styles["Normal"].font.name = "Calibri"
    d.add_heading("All projects at a glance", 0)
    d.add_paragraph("One row per project. Each has its own detailed guide with architecture in this folder.")
    t = d.add_table(rows=1, cols=3)
    t.style = "Light Grid Accent 1"
    for c, txt in zip(t.rows[0].cells, ["Project", "What it does", "Main technology"], strict=True):
        c.text = txt
    for p in PROJECTS:
        row = t.add_row().cells
        row[0].text, row[1].text, row[2].text = p["name"], p["tagline"], p["stack"]
    out = OUT / "00_All_Projects_Overview.docx"
    d.save(out)
    return out


if __name__ == "__main__":
    for proj in PROJECTS:
        print(build(proj))
    print(overview())
