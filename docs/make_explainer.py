"""Builds docs/EvalKit_Explained.docx, a plain-language guide to EvalKit. Run: python docs/make_explainer.py"""
# ruff: noqa: E501

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

d = Document()
d.styles["Normal"].font.name = "Calibri"
d.styles["Normal"].font.size = Pt(11)
for s in d.sections:
    s.left_margin = s.right_margin = Inches(1)
    s.top_margin = s.bottom_margin = Inches(0.9)


def h(text, lvl=1):
    d.add_heading(text, lvl)


def p(text, bold=False, italic=False):
    r = d.add_paragraph().add_run(text)
    r.bold, r.italic = bold, italic


def bl(items):
    for i in items:
        para = d.add_paragraph(style="List Bullet")
        if isinstance(i, tuple):
            para.add_run(i[0]).bold = True
            para.add_run(" " + i[1])
        else:
            para.add_run(i)


def table(rows, widths):
    t = d.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Light Grid Accent 1"
    for i, row in enumerate(rows):
        for j, c in enumerate(row):
            cell = t.cell(i, j)
            cell.width = Inches(widths[j])
            r = cell.paragraphs[0].add_run(c)
            r.font.size = Pt(10)
            r.bold = i == 0
    d.add_paragraph()


d.add_heading("EvalKit", 0)
p("A plug-in toolkit that tells you, with numbers, whether your AI app is actually good", italic=True)
p("A plain-language guide: what it is, why it matters, what it can do, and how to explain it to others.")

h("1. The one-minute version")
p("You build an AI app: a chatbot, a document search, a resume screener. It gives answers. But how do you know the "
  "answers are any good? Most people try a few questions, say \"looks fine\", and ship it. Then a small change "
  "breaks something and nobody notices.")
p("EvalKit is a test machine for AI apps. You give it a list of questions with the right answers. It asks your app "
  "every question, checks each reply, and shows you a score, the questions that failed, and whether your latest "
  "change made things better or worse.")
p("Think of it as school exams for your AI: the same exam every time, graded automatically, with a report card you "
  "can compare week to week.", bold=True)

h("2. The problem it solves")
table([
    ["Without EvalKit", "With EvalKit"],
    ["\"It seems to work\" after trying 5 questions", "A score across dozens or hundreds of questions"],
    ["A change fixes one thing and quietly breaks another", "Every run is compared with the last; regressions are flagged"],
    ["No idea if a new model or prompt is really better", "A statistical test says whether the gap is real or just luck"],
    ["Bad answers found by angry users", "Bad answers found before release, automatically in CI"],
    ["No idea what it costs or how slow it is", "Cost and speed measured on every question"],
    ["A new custom script for every app", "Connect any app in minutes, no code"],
], [3.2, 3.3])

h("3. What it can do")
bl([
    ("Test any AI app.", "Chatbots, RAG (document question answering), agents, or model APIs such as OpenAI, Claude, "
     "Gemini, Ollama and Groq."),
    ("Connect without writing code.", "Paste a cURL command, pick a service from a list, or give a web address. EvalKit "
     "sends one test question, works out where the answer is in the reply, and writes the settings for you."),
    ("Handle awkward apps.", "Streaming replies, file uploads, websites with a login and HTML pages, and plain "
     "Python functions."),
    ("Grade answers many ways.", "Exact match, contains-text, pattern match, meaning similarity, whether the answer "
     "sticks to the retrieved documents, whether the app says \"I don't know\" when it should, speed and cost "
     "limits, and an AI judge that scores answers 1 to 5."),
    ("Check search quality.", "For document-search apps: was the right document returned (hit rate), how high was it "
     "ranked (MRR), how much was found (recall)."),
    ("Compare two versions fairly.", "Paired statistical tests say whether version B is truly better than A, or the "
     "gap is noise."),
    ("Block bad releases.", "A quality gate gives pass or fail for CI: for example, fail the build if accuracy "
     "drops more than 5%."),
    ("Generate tests.", "Create test questions from your documents and run built-in red-team attacks "
     "(prompt injection, secret leaking and more)."),
    ("Check consistency.", "Ask the same question many times, or rephrased, and see if answers stay stable."),
    ("Show a dashboard.", "Overview, run history, per-question detail, side-by-side comparison, live progress."),
    ("Make safe copies of data.", "The snapshot command copies a database or folder so tests never touch real data."),
])

h("4. How it works, step by step")
bl([
    ("Step 1, Connect.", "Tell EvalKit how to reach your app, then click Test connection."),
    ("Step 2, Test set.", "A file of questions plus the expected answer or the right document. "
     "20 to 50 good questions is a great start."),
    ("Step 3, Choose scorers.", "Pick how answers are checked. You can mix several."),
    ("Step 4, Run.", "EvalKit asks every question, in parallel, and grades each reply."),
    ("Step 5, Read the report.", "Pass rate with a confidence range, failures with the reason, speed, cost, "
     "results by topic."),
    ("Step 6, Change and compare.", "Tweak your app, run again, and see if it really improved."),
    ("Step 7, Automate.", "Put it in CI so every code change is tested before release."),
])

h("5. Who it helps")
table([
    ["Person", "What they get"],
    ["Developer", "Confidence to change prompts, models and settings without silent breakage"],
    ["Team lead or manager", "One number and a trend line instead of opinions, plus a release gate"],
    ["Product owner", "Evidence that quality is improving, with cost and speed figures"],
    ["Recruiter or interviewer", "Proof the builder knows how to test AI, which is the hard, valued part of the job"],
], [2.0, 4.5])

h("6. Real results from this project")
p("EvalKit was tried on apps in the PROJECTS folder, with no custom code per app:")
table([
    ["App", "How it connected", "Result"],
    ["incident-copilot (agent)", "Fresh conversation id per question", "4 of 4 passed"],
    ["PaperGraph (knowledge graph chat)", "Opens a chat session, then asks",
     "2 of 6 passed. Found it answers out-of-scope questions with junk and had a 57 s cold start"],
    ["finqa-agent (retrieval)", "Graded the retrieval results", "hit@5 0.92, MRR 0.86 on 12 questions"],
    ["rag-chatbot (streaming)", "Reads streamed replies",
     "The app's own Gemini quota was used up; EvalKit reported it clearly as an error"],
], [2.1, 2.1, 2.3])
p("Along the way the tests exposed and fixed four bugs in EvalKit itself, which is the toolkit doing its job.")

h("7. Honest limits")
bl([
    "Best for question-and-answer and document-search apps. Not built for image, video, voice or scraper apps.",
    "A keyword check is crude. For real quality judgment use the AI judge, which needs an LLM key "
    "(the free Gemini tier allows about 20 requests a day per model).",
    "Results are only as good as the test set. A tiny test set gives wide confidence ranges.",
    "The website-login feature was tested on a simulated site, not yet on a real one such as the resume screener.",
    "It has been tried on the author's own projects only, not by outside users yet.",
])

h("8. Words explained simply")
table([
    ["Word", "Meaning"],
    ["Test set", "A list of questions with the right answers, like an exam paper"],
    ["Scorer", "The marker that grades one answer"],
    ["Pass rate", "Percent of questions answered correctly"],
    ["Confidence range", "How sure we are about that percent; small test sets give wide ranges"],
    ["RAG", "An AI app that first looks up documents, then answers from them"],
    ["Hit rate, MRR, recall", "Scores for how well the search step found the right documents"],
    ["LLM judge", "Another AI that reads an answer and grades it 1 to 5"],
    ["A/B test", "Comparing version A with version B on the same questions"],
    ["Regression", "Something that used to work and now does not"],
    ["Quality gate", "A pass or fail rule that can stop a bad release"],
    ["Red-team", "Attack-style questions that try to trick the app"],
    ["Snapshot", "A safe copy of data so tests never touch the real thing"],
], [1.8, 4.7])

h("9. How to explain it in 30 seconds")
p("\"AI apps are hard to test because answers are text, not true or false. I built EvalKit. It plugs into any AI app "
  "in minutes, runs a set of questions, scores the answers, and tells you statistically whether a change made "
  "things better or worse. It can block a bad release in CI, and I used it on my own apps to find real problems.\"",
  italic=True)

h("10. How to run it")
bl([
    "Start everything: docker compose up -d --build, then open http://localhost:8800.",
    "Optional Gemini key: put it in the .env file, never in code or on GitHub. Free at aistudio.google.com/apikey.",
    "Command line: evalkit connect, evalkit run, evalkit compare, evalkit snapshot.",
    "Ready-made configs for sample apps are in examples/own_projects/.",
])

h("Screenshots")
for f, cap in [("overview", "Overview: pass rate, trend, recent runs"),
               ("compare", "Compare: did the change really help?"),
               ("run-detail", "Run detail: every question and why it failed")]:
    d.add_picture(f"docs/screenshots/{f}.jpg", width=Inches(6.2))
    c = d.add_paragraph(cap)
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.runs[0].italic = True

d.save("docs/EvalKit_Explained.docx")
